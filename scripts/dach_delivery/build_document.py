#!/usr/bin/env python3
"""Build the formal DOCX from the reviewed Markdown and a vector-derived figure."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs'
FONT = 'Arial'
NAVY, BLUE, MUTED = '183047', '286C9E', '536778'

def shade(cell, fill):
    el=OxmlElement('w:shd'); el.set(qn('w:fill'),fill); cell._tc.get_or_add_tcPr().append(el)

def inline(p, s, table=False):
    for piece in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)',s):
        r=p.add_run(piece[1:-1] if piece.startswith('`') else piece[2:-2] if piece.startswith('**') else piece)
        if piece.startswith('`'):
            r.font.name='Liberation Mono'; r.font.size=Pt(8.0 if table else 8.6)
            r.font.color.rgb=RGBColor.from_string(NAVY)
        elif piece.startswith('**'):r.bold=True
    return p

def page_setup(section,landscape=False):
    section.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width=Mm(297 if landscape else 210); section.page_height=Mm(210 if landscape else 297)
    section.left_margin=section.right_margin=Mm(15 if landscape else 18)
    section.top_margin=Mm(13 if landscape else 15); section.bottom_margin=Mm(12 if landscape else 15)
    section.footer_distance=Mm(7)

def bookmark(p, number):
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(number)); start.set(qn('w:name'),'s'+p.text.split()[0].replace('.','_'))
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number));p._p.insert(0,start);p._p.append(end)

def cells(line):
    return [s.strip() for s in re.split(r'\|(?=(?:[^`]*`[^`]*`)*[^`]*$)',line.strip().strip('|'))]

def main():
    doc=Document();page_setup(doc.sections[0]);doc.core_properties.title='DACH catalogue ingestion specification'
    doc.core_properties.author='ARES'; doc.core_properties.subject='UMAS v3 content and counters ingestion contracts'
    styles=doc.styles

    for st in styles:
        if hasattr(st, '_element'):
            for border in st._element.xpath('.//w:pBdr'):
                border.getparent().remove(border)
    normal=styles['Normal'];normal.font.name=FONT;normal.font.size=Pt(10)
    normal.font.color.rgb=RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.04
    normal.paragraph_format.widow_control=True
    for name,size,color in [('Title',24,'000000'),('Heading 1',15,'000000'),('Heading 2',11,'000000')]:
        st=styles[name];st.font.name=FONT;st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color)
        st.font.bold=True;st.paragraph_format.space_before=Pt(10 if name!='Title' else 0)
        st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
    styles['List Bullet'].font.name=FONT;styles['List Bullet'].font.size=Pt(10)
    styles['List Bullet'].paragraph_format.space_after=Pt(5)
    styles['List Bullet'].paragraph_format.line_spacing=1.04
    footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('ARES  /  DACH     ');r.font.name=FONT;r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(MUTED)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    lines=(ROOT/'scripts/dach_delivery/document.md').read_text().splitlines();i=0;heading_id=0
    while i<len(lines):
        ln=lines[i];i+=1
        if not ln.strip():continue
        if ln.startswith('<!-- page '):
            landscape='landscape' in ln
            sec=doc.add_section(WD_SECTION.NEW_PAGE);page_setup(sec,landscape)
        elif ln.startswith('<!--'):continue
        elif ln=='[FIGURE]':
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0)
            r=p.add_run();r.add_picture(str(OUT/'figures/ingestion-three-lanes.png'),width=Mm(267))
            r._r.xpath('.//wp:docPr')[0].set('descr','Three source-owned ingestion lanes converge on one content index, with external Metadata reject notification through AWS SQS.')
        elif ln.startswith('# '):doc.add_paragraph(ln[2:],'Title')
        elif ln.startswith('## '):
            p=doc.add_paragraph(ln[3:],'Heading 1');heading_id+=1;bookmark(p,heading_id)
        elif ln.startswith('### '):
            p=doc.add_paragraph(ln[4:],'Heading 2');heading_id+=1;bookmark(p,heading_id)
        elif ln.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),ln[2:])
        elif ln.startswith('```'):
            code=[]
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            i+=1
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(7)
            p.paragraph_format.keep_together=True
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F5F8');p._p.get_or_add_pPr().append(sh)
            r=p.add_run('\n'.join(code));r.font.name='Liberation Mono';r.font.size=Pt(8.4)
        elif ln.startswith('|'):
            rows=[cells(ln)]
            while i<len(lines) and lines[i].startswith('|'):
                v=cells(lines[i]);i+=1
                if all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in v):continue
                rows.append(v)
            table=doc.add_table(rows=0,cols=len(rows[0]));table.autofit=False
            borders=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
            table._tbl.tblPr.append(borders)
            widths={2:[.31,.69],3:[.32,.26,.42],4:[.285,.175,.13,.41]}[len(rows[0])]
            dictionary=rows[0][0]=='Index field'
            if rows[0][0]=='Artifact':widths=[.47,.53]
            usable=(doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin)/914400
            for c,w in zip(table.columns,widths):c.width=Inches(usable*w)
            for n,rowdata in enumerate(rows):
                row=table.add_row();pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
                if n==0:head=OxmlElement('w:tblHeader');pr.append(head)
                if len(rowdata)!=len(rows[0]):raise ValueError('Inconsistent table columns: '+str(rowdata))
                for j,txt in enumerate(rowdata):
                    c=row.cells[j];c.width=Inches(usable*widths[j]);p=c.paragraphs[0]
                    p.paragraph_format.space_before=Pt(1 if dictionary else 3);p.paragraph_format.space_after=Pt(1 if dictionary else 4)
                    p.paragraph_format.line_spacing=1.02
                    inline(p,txt,table=True)
                    for r in p.runs:
                        if r.font.name!='Liberation Mono':r.font.size=Pt(9)
                        if n==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
                    shade(c,NAVY if n==0 else 'F1F5F8' if n%2 else 'FAFBFC')
                    margins=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        m=OxmlElement('w:'+side);m.set(qn('w:w'),'45' if dictionary and side in ['top','bottom'] else '85');m.set(qn('w:type'),'dxa');margins.append(m)
                    c._tc.get_or_add_tcPr().append(margins)
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.add_run().font.size=Pt(3)
        else:inline(doc.add_paragraph(),ln)
    OUT.mkdir(exist_ok=True)
    target=OUT/'ARES-DACH-Ingestion-Specification.docx';doc.save(target);print(target)

if __name__=='__main__':main()
